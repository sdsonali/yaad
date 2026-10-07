"""10-slide NL_GooglePhotos. 16:9, 14pt, no fellow name. python survey/build_deck.py"""

from __future__ import annotations

import json
import os
import sys
from pathlib import Path

os.environ.setdefault("PYTHONIOENCODING", "utf-8")

_ROOT = Path(__file__).resolve().parents[1]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from shared import config

OUT = _ROOT / "NL_GooglePhotos.pptx"
CHART = _ROOT / "deck" / "assets" / "funnel_stages.png"
LIVE = "https://REPLACE_ME.streamlit.app/"
NOTES = "survey/interview_notes.md"

SLATE = RGBColor(0x1B, 0x2A, 0x4A)
TEAL = RGBColor(0x2F, 0x6F, 0x7A)
BODY = RGBColor(0x1F, 0x24, 0x30)
MUTED = RGBColor(0x5C, 0x65, 0x70)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SAND = RGBColor(0xED, 0xE6, 0xD9)
MIST = RGBColor(0xE4, 0xEE, 0xF0)
PAGE = RGBColor(0xF7, 0xF5, 0xF2)
W, H = Inches(13.333), Inches(7.5)


def run(p, size=14, bold=False, color=BODY):
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = "Calibri"


def tb(slide, l, t, w, h, text, size=14, bold=False, color=BODY):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = text
    run(r, size, bold, color)
    return box


def fill(shape, rgb):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb
    shape.line.fill.background()


def bg(slide):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    fill(sh, PAGE)


def panel(slide, l, t, w, h, col):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    fill(sh, col)


def footer(slide, n):
    tb(slide, 0.5, 7.15, 10, 0.28, "Product Manager, Core Experience", 14, False, MUTED)
    tb(slide, 12.1, 7.15, 0.8, 0.28, str(n), 14, False, MUTED)


def main() -> None:
    opp = json.loads(config.OPPORTUNITY_JSON.read_text(encoding="utf-8")) if config.OPPORTUNITY_JSON.exists() else {}
    test = json.loads(config.MVP_TEST_JSON.read_text(encoding="utf-8")) if config.MVP_TEST_JSON.exists() else {}
    stages = (opp.get("top_funnel_stages") or [])[:6]
    if stages:
        fig, ax = plt.subplots(figsize=(8.2, 3.0))
        ax.bar([s["label"].replace("_", "\n") for s in stages], [s["pct_of_mentions"] for s in stages], color="#1B2A4A")
        ax.set_ylabel("% of tagged items")
        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)
        CHART.parent.mkdir(parents=True, exist_ok=True)
        fig.tight_layout()
        fig.savefig(CHART, dpi=130, facecolor="#F7F5F2")
        plt.close()

    n = opp.get("total_reviews") or 0
    headline = opp.get("headline") or "Funnel stages ranked from gold corpus."
    found, miss = test.get("found"), test.get("miss")

    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    bg(s)
    tb(s, 0.5, 0.35, 12, 0.3, "PRODUCT STRATEGY — GOOGLE PHOTOS · CORE EXPERIENCE", 14, True, TEAL)
    tb(s, 0.5, 0.75, 12.3, 1.3, "TOPIC SEARCH RETURNS ONE PILE — PHOTOS, PAPERS, SCREENSHOTS", 26, True, SLATE)
    tb(s, 0.5, 2.2, 12.3, 1.0, "People search a life topic they still remember. Photos answers with one undifferentiated grid. The intended item can sit in the set while they evaluate the wrong kind.", 16)
    tb(s, 0.5, 3.3, 12.3, 0.7, "Raise successful evaluation of topic-search results by separating personal photos, documents, and screenshots — not by asking for a better keyword.", 16, True, SLATE)
    for i, (k, v) in enumerate([("FOCUS", "Coarse life-topic search"), ("LEVER", "Type mix / visual evaluation"), ("OUT", "Query rewrite, delete, father-only People")]):
        panel(s, 0.5 + i * 4.2, 4.3, 4.0, 1.5, SAND if i != 1 else MIST)
        tb(s, 0.7 + i * 4.2, 4.45, 3.6, 0.35, k, 14, True, TEAL)
        tb(s, 0.7 + i * 4.2, 4.9, 3.6, 0.7, v)
    tb(s, 0.5, 6.05, 12, 0.4, f"Live app {LIVE}  ·  Notes {NOTES}", 14, True, TEAL)
    footer(s, 1)

    s = prs.slides.add_slide(blank)
    bg(s)
    tb(s, 0.5, 0.3, 12.3, 1.0, "THE LEAK IS EVALUATION — THE ITEM MAY ALREADY BE IN THE SET", 22, True, SLATE)
    tb(s, 0.5, 1.3, 12.3, 0.55, "North star: share of topic-search sessions where the user can confirm the intended kind of item.", 14)
    stages_n = ["Recall", "Query", "Interpret", "RELEVANCE", "EVALUATE", "Refine"]
    for i, name in enumerate(stages_n):
        hot = name in ("RELEVANCE", "EVALUATE")
        panel(s, 0.4 + i * 2.15, 2.05, 2.05, 0.8, SLATE if hot else TEAL)
        tb(s, 0.45 + i * 2.15, 2.2, 1.95, 0.5, name, 14, True, WHITE)
    lines = [
        "Recall — they have a topic: college, a trip, “that certificate.”",
        "Query formulation — they can type the topic. Not the hero failure.",
        "Relevance — search mixes camera photos, screenshots, and documents that share the topic word. TARGET.",
        "Evaluate — certificates look like “college”; they cannot pick the memory. TARGET.",
        "Refine — they scroll, open Files, or leave. Intelligence belongs in typing the grid, not rewriting the box.",
        "IN: large libraries, topic search. OUT: crash, storage, People-tab only, auto-delete.",
    ]
    for i, line in enumerate(lines):
        tb(s, 0.5, 3.05 + i * 0.62, 12.3, 0.6, line)
    footer(s, 2)

    s = prs.slides.add_slide(blank)
    bg(s)
    tb(s, 0.5, 0.3, 12.3, 0.9, "THE ENGINE COMPARES FUNNEL FAILURES — IT DOES NOT SCORE SENTIMENT", 22, True, SLATE)
    panel(s, 0.5, 1.4, 6.1, 5.0, MIST)
    tb(s, 0.7, 1.55, 5.7, 0.35, "PIPELINE", 14, True, TEAL)
    tb(s, 0.7, 2.05, 5.7, 4.0, "Play · App Store · Reddit · YouTube\n→ filter: retrieval + topic / mixed-kind talk (drop crash/storage-only)\n→ tag: funnel stage, memory type, anchors, workaround, severity\n→ matrix for the deck\n→ live Q&A over tagged items (TF-IDF + counts)\n\nWorked question: “when do certificates show up in photo search?”")
    panel(s, 6.9, 1.4, 5.9, 5.0, SAND)
    tb(s, 7.1, 1.55, 5.5, 0.35, "CLICK PATH", 14, True, TEAL)
    tb(s, 7.1, 2.05, 5.5, 4.0, f"1. Open live app → Discovery.\n2. Theme chip or type a PM question.\n3. Read findings tied to tagged fields.\n4. Compare funnel × memory-type table.\n\nGold n = {n}.\nSources: {opp.get('by_source')}\n{headline}")
    footer(s, 3)

    s = prs.slides.add_slide(blank)
    bg(s)
    tb(s, 0.5, 0.3, 12.3, 0.8, "CORPUS RANKS STAGES — INTERVIEWS LOCK THE TYPE-MIX JOB", 20, True, SLATE)
    tb(s, 0.5, 1.15, 12.3, 0.45, headline)
    if CHART.exists():
        s.shapes.add_picture(str(CHART), Inches(0.4), Inches(1.7), Inches(8.0), Inches(3.4))
    types = opp.get("top_memory_types") or []
    ttxt = "\n".join(f"{r['label']}: {r['pct_of_mentions']}%" for r in types[:5]) or "from live engine"
    tb(s, 8.6, 1.7, 4.3, 4.5, "Memory types\n\n" + ttxt + "\n\nHero IN: type-blind topic grids.\nCaveat: reviews under-index silent scroll. Interviews are a second layer — not averaged. No invented lift %.")
    footer(s, 4)

    s = prs.slides.add_slide(blank)
    bg(s)
    tb(s, 0.5, 0.3, 12.3, 0.8, "RITU AND SWATI LOCK THE JOB — RAHUL AND AKSHAY ARE NAMED OUT", 22, True, SLATE)
    cards = [
        ("Ritu — IN", "Searched college photos. The same grid returned certificates. She could not evaluate which tiles were memories."),
        ("Swati — IN", "Wants personal photos, documents, and screenshots sequenced or separated — not one mixed dump."),
        ("Named OUT", "Rahul: father-only People search (not generic search; do not demo a deceased parent). Akshay: auto-delete unused photos — cleanup, off metric."),
    ]
    for i, (t, b) in enumerate(cards):
        panel(s, 0.5 + i * 4.2, 1.4, 4.0, 2.7, MIST if i != 1 else SAND)
        tb(s, 0.7 + i * 4.2, 1.55, 3.6, 0.45, t, 14, True, TEAL)
        tb(s, 0.7 + i * 4.2, 2.1, 3.6, 1.8, b)
    tb(s, 0.5, 4.35, 12.3, 2.0, "Method: 30–40 min critical-incident reconstruction. First names only. Directional n. Full notes in the repo. Two evidence layers — corpus % is not interview n.", 14)
    footer(s, 5)

    s = prs.slides.add_slide(blank)
    bg(s)
    tb(s, 0.5, 0.3, 12.3, 0.85, "RESEARCH CONFIRMS TYPE-BLIND GRIDS — AND KILLS THE WRONG FIXES", 22, True, SLATE)
    panel(s, 0.5, 1.35, 6.0, 5.1, MIST)
    tb(s, 0.7, 1.5, 5.6, 0.35, "CONFIRMED", 14, True, TEAL)
    tb(s, 0.7, 2.05, 5.6, 4.0, "• They can name the topic.\n• Photos mixes memories with papers and screenshots.\n• The right tile can be present and still unusable.\n• They want kinds sequenced or separated.\n• Workarounds: scroll, Files, leave.")
    panel(s, 6.8, 1.35, 6.0, 5.1, SAND)
    tb(s, 7.0, 1.5, 5.6, 0.35, "KILLED", 14, True, TEAL)
    tb(s, 7.0, 2.05, 5.6, 4.0, "• Killed as hero: query formulation / smarter keywords.\n• Killed: date picker as clarifying UI.\n• Killed: Rahul — People-tab / father-only.\n• Killed: Akshay — auto-delete unused photos.\n• Corpus % ≠ interview n. Never merged. No fake lift.")
    footer(s, 6)

    s = prs.slides.add_slide(blank)
    bg(s)
    tb(s, 0.5, 0.3, 12.3, 0.8, "ONE GRID CANNOT CARRY THREE KINDS OF “COLLEGE”", 22, True, SLATE)
    cells = [
        ("SEGMENT", "Large, multi-year libraries. Topic search, not filename search."),
        ("SCENARIO", "“College photos” returns memories + certificates + screenshots together."),
        ("OUTCOME", "Visual evaluation + type-relevant ranking."),
        ("ROOT CAUSE", "Topic words are shared across kinds. The product does not type the grid."),
        ("WORKAROUNDS", "Scroll past papers, open Files, ask someone."),
        ("WHY IT MATTERS", "Topic search stays the default. Photos remains the archive, not chat apps."),
    ]
    for i, (k, v) in enumerate(cells):
        r, c = divmod(i, 3)
        panel(s, 0.45 + c * 4.25, 1.25 + r * 2.4, 4.1, 2.2, SAND if c != 1 else MIST)
        tb(s, 0.6 + c * 4.25, 1.4 + r * 2.4, 3.8, 0.35, k, 14, True, TEAL)
        tb(s, 0.6 + c * 4.25, 1.85 + r * 2.4, 3.8, 1.4, v)
    tb(s, 0.5, 6.15, 12.3, 0.6, "Metric → funnel (relevance + evaluate) → Ritu/Swati → this problem. Not “search is hard.”", 14, True, SLATE)
    footer(s, 7)

    s = prs.slides.add_slide(blank)
    bg(s)
    tb(s, 0.5, 0.3, 12.3, 0.75, "ASK KIND ONCE — THEN SHOW A TYPED GRID. NOT A CHATBOT.", 22, True, SLATE)
    panel(s, 0.45, 1.2, 6.3, 5.2, MIST)
    tb(s, 0.65, 1.35, 6.0, 0.35, "MVP LOOP", 14, True, TEAL)
    tb(s, 0.65, 1.8, 5.9, 4.3, "1. Dump a coarse topic (college photos).\n2. If results mix kinds, ask: memory vs saved to use later (document/screenshot).\n3. Cap at two turns (optional document vs screenshot).\n4. Show a typed thumbnail grid: personal / document / screenshot.\n5. None of these → different kind.\n\nCollege mixed-results card is the hero task. Demo library only.")
    panel(s, 6.95, 1.2, 5.9, 5.2, SAND)
    tb(s, 7.15, 1.35, 5.5, 0.35, "PART 6", 14, True, TEAL)
    tb(s, 7.15, 1.8, 5.5, 4.3, f"n=3 prompt cards (college memories, college certificate, saved screenshot).\nFound in top 5: {found}  miss: {miss}\n\nLog: original query, clarifying turns, success/fail.\nNo lift % invented.\n\nNot built: Google login, date picker, open chat, auto-delete, father-only filter.")
    footer(s, 8)

    s = prs.slides.add_slide(blank)
    bg(s)
    tb(s, 0.5, 0.3, 12.3, 0.75, "MEASURE WHETHER THEY CAN EVALUATE KIND — NOT QUERY REWRITE", 22, True, SLATE)
    metrics = [
        "North star — Topic-search evaluation success: in-scope sessions that end in a confirmed find of the intended kind.",
        "Leading — Type-mix ask rate: share of sessions whose first grid mixed personal / document / screenshot.",
        "Leading — Success after kind answer: memory vs saved-for-later actually changes the usable set.",
        "Diagnostic — Wrong-kind selects (certificate when they wanted photos).",
        "Diagnostic — Question rated helpful vs annoying (Part 6 free text).",
        "Guardrail — Precise keyword search must not regress; no auto-delete; demo library only (privacy).",
    ]
    for i, line in enumerate(metrics):
        tb(s, 0.5, 1.2 + i * 0.85, 12.3, 0.8, line)
    footer(s, 9)

    s = prs.slides.add_slide(blank)
    bg(s)
    tb(s, 0.5, 0.3, 12.3, 0.75, "A TYPED DEMO GRID CAN LOOK SOLVED AND STILL FAIL IN PHOTOS", 22, True, SLATE)
    risks = [
        "Demo cheating — small set. Mitigation: keep utility library + college memories and certificates that share topic words; report miss rate.",
        "Date-shaped questions sneak back — wrong lock. Mitigation: scripted type-mix questions only; no date picker.",
        "Looks like ChatGPT-on-search — Creativity collapse. Mitigation: two-turn cap, no open chat, typed grid is the UI.",
        "Scope creep to People or delete — Rahul/Akshay. Mitigation: named OUT in docs and deck.",
        "Privacy — testers will not use a real account. Mitigation: bundled demo library; production stays inside Photos.",
    ]
    for i, line in enumerate(risks):
        tb(s, 0.5, 1.2 + i * 1.05, 12.3, 1.0, line)
    footer(s, 10)

    prs.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
